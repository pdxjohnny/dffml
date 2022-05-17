from pptx import Presentation
from pptx.util import Inches


DATA = [
    {
        "key": "https://github.com/intel/dffml",
        "authors": [
            6,
            5,
            9,
            10,
            16,
            8,
            10,
            22,
            15,
            10
        ],
        "commits": [
            33,
            58,
            327,
            58,
            180,
            191,
            129,
            277,
            243,
            66
        ],
        "work": [
            37,
            40,
            75,
            77,
            45,
            10,
            80,
            82,
            68,
            22
        ]
    }
]


def gen_pptx():
    # --- Picture

    img_path = 'monty-truth.png'

    import matplotlib.pyplot as plt

    repo = DATA[0]
    key = repo["key"]

    features = repo.copy()
    del features["key"]

    for feature, data in features.items():
        if isinstance(data, list):
            plt.plot(data)
            plt.ylabel(feature)
            plt.show()
            plt.savefig(f'{feature}-' + img_path)


    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    blank_slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_slide_layout)

    left = top = Inches(1)
    pic = slide.shapes.add_picture(img_path, left, top)

    left = Inches(5)
    height = Inches(5.5)
    pic = slide.shapes.add_picture(img_path, left, top, height=height)


    # --- Table

    prs = Presentation()
    title_only_slide_layout = prs.slide_layouts[5]
    slide = prs.slides.add_slide(title_only_slide_layout)
    shapes = slide.shapes

    shapes.title.text = 'Adding a Table'

    rows = cols = 2
    left = top = Inches(2.0)
    width = Inches(6.0)
    height = Inches(0.8)

    table = shapes.add_table(rows, cols, left, top, width, height).table

    # set column widths
    table.columns[0].width = Inches(2.0)
    table.columns[1].width = Inches(4.0)

    # write column headings
    table.cell(0, 0).text = 'Foo'
    table.cell(0, 1).text = 'Bar'

    # write body cells
    table.cell(1, 0).text = 'Baz'
    table.cell(1, 1).text = 'Qux'

    prs.save('test.pptx')
